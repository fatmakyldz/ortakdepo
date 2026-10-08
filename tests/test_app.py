import io
from datetime import date, timedelta

from openpyxl import load_workbook
from PIL import Image

from app.config import settings
from app.models import Category, FixedExpense, Installment, Receipt, Settlement, Transaction, User
from app.timeutil import today


def _png(color=(200, 30, 30), size=(3000, 1200)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def _cat(db, name):
    return db.query(Category).filter_by(name=name).one().id


def test_redirects_to_setup_then_login(client):
    assert client.get("/", follow_redirects=False).headers["location"] == "/kurulum"
    assert client.get("/kurulum").status_code == 200


def test_setup_validation_and_single_use(client, db):
    r = client.post("/kurulum", data={"ad1": "F", "eposta1": "kotu", "sifre1": "kisa",
                                      "ad2": "", "eposta2": "k@x.co", "sifre2": "yeterince-uzun"})
    assert r.status_code == 422 and db.query(User).count() == 0
    ok = {"ad1": "Fatma", "eposta1": "f@x.co", "sifre1": "gizli-sifre-1",
          "ad2": "Kerem", "eposta2": "k@x.co", "sifre2": "gizli-sifre-2"}
    assert client.post("/kurulum", data=ok, follow_redirects=False).status_code == 303
    # ikinci kez kurulum yapılamaz
    client.post("/kurulum", data={**ok, "eposta1": "x@x.co", "eposta2": "y@x.co"})
    assert db.query(User).count() == 2


def test_login_logout_and_throttle(logged_in):
    c = logged_in
    c.post("/cikis")
    assert c.get("/", follow_redirects=False).headers["location"] == "/giris"
    assert c.post("/giris", data={"eposta": "fatma@example.com", "sifre": "yanlis"}).status_code == 401
    r = c.post("/giris", data={"eposta": "FATMA@example.com ", "sifre": "gizli-sifre-1"}, follow_redirects=False)
    assert r.status_code == 303
    c.post("/cikis")
    for _ in range(8):
        c.post("/giris", data={"eposta": "fatma@example.com", "sifre": "yanlis"})
    assert c.post("/giris", data={"eposta": "fatma@example.com", "sifre": "gizli-sifre-1"}).status_code == 429


def test_cross_site_post_is_rejected(logged_in):
    r = logged_in.post("/cikis", headers={"origin": "https://kotu.example"})
    assert r.status_code == 403
    assert logged_in.get("/").status_code == 200  # oturum hâlâ açık


def test_pages_require_login(client, db):
    client.post("/kurulum", data={"ad1": "F", "eposta1": "f@x.co", "sifre1": "gizli-sifre-1",
                                  "ad2": "K", "eposta2": "k@x.co", "sifre2": "gizli-sifre-2"})
    client.post("/cikis")
    for url in ["/", "/kayitlar", "/sabit", "/ortaklar", "/rapor", "/ayarlar", "/rapor/excel?bas=2026-01-01&bit=2026-12-31", "/fisler/1/dosya"]:
        r = client.get(url, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == "/giris", url


def test_expense_with_receipt_roundtrip(logged_in, db):
    c = logged_in
    fatma = db.query(User).filter_by(name="Fatma").one()
    r = c.post("/kayitlar/yeni", data={
        "tur": "gider", "gun": today().isoformat(), "tutar": "1.250,50",
        "kategori_id": str(_cat(db, "Yakıt")), "kim": str(fatma.id), "aciklama": "34 ABC 123 depo",
    }, files=[("fisler", ("fis.png", _png(), "image/png")), ("fisler", ("fatura.pdf", b"%PDF-1.4 test", "application/pdf"))],
        follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/"
    tx = db.query(Transaction).one()
    assert (tx.amount, tx.kind, tx.partner_id, tx.created_by_id) == (125050, "gider", fatma.id, fatma.id)
    img, pdf = tx.receipts
    assert img.content_type == "image/jpeg" and pdf.content_type == "application/pdf"

    # fotoğraf küçültülmüş ve küçük resmi üretilmiş olmalı
    full = settings.upload_dir / img.path
    assert max(Image.open(full).size) == 2000
    assert max(Image.open(settings.upload_dir / img.thumb_path).size) == 480

    assert c.get(f"/fisler/{img.id}/dosya").headers["content-type"] == "image/jpeg"
    assert c.get(f"/fisler/{img.id}/kucuk").status_code == 200
    assert c.get(f"/fisler/{pdf.id}/dosya").content == b"%PDF-1.4 test"

    page = c.get("/").text
    assert "Yakıt" in page and "1.250,50" in page
    assert "34 ABC 123 depo" in c.get("/kayitlar").text
    assert c.get(f"/kayitlar/{tx.id}").status_code == 200

    # silince dosyalar da diskten kalkar
    c.post(f"/kayitlar/{tx.id}/sil")
    assert db.query(Transaction).count() == 0 and db.query(Receipt).count() == 0
    assert not full.exists()


def test_invalid_entries_are_rejected(logged_in, db):
    c = logged_in
    base = {"tur": "gider", "gun": today().isoformat(), "tutar": "100",
            "kategori_id": str(_cat(db, "Yemek")), "kim": "ortak"}
    for patch in [{"tutar": ""}, {"tutar": "abc"}, {"tutar": "0"}, {"kategori_id": ""},
                  {"kategori_id": str(_cat(db, "İş geliri"))},           # gelir kategorisi gidere yazılamaz
                  {"kim": "999"}, {"kim": ""}, {"gun": ""},
                  {"gun": (today() + timedelta(days=5)).isoformat()}]:
        assert c.post("/kayitlar/yeni", data={**base, **patch}).status_code == 422, patch
    assert db.query(Transaction).count() == 0
    # bozuk dosya kaydı engellemez ama uyarı verir
    r = c.post("/kayitlar/yeni", data=base, files=[("fisler", ("x.jpg", b"resim degil", "image/jpeg"))])
    assert db.query(Transaction).count() == 1 and db.query(Receipt).count() == 0
    assert "okunamadı" in r.text


def test_edit_keeps_receipts_and_updates_fields(logged_in, db):
    c = logged_in
    c.post("/kayitlar/yeni", data={"tur": "gelir", "gun": today().isoformat(), "tutar": "5000",
                                   "kategori_id": str(_cat(db, "İş geliri")), "kim": "ortak"},
           files=[("fisler", ("d.png", _png(size=(50, 50)), "image/png"))])
    tx = db.query(Transaction).one()
    kerem = db.query(User).filter_by(name="Kerem").one()
    c.post(f"/kayitlar/{tx.id}/duzenle", data={"gun": today().isoformat(), "tutar": "5.500,00",
                                               "kategori_id": str(_cat(db, "Diğer gelir")),
                                               "kim": str(kerem.id), "aciklama": "düzeltildi"})
    db.refresh(tx)
    assert (tx.amount, tx.partner_id, tx.description, tx.kind) == (550000, kerem.id, "düzeltildi", "gelir")
    assert len(tx.receipts) == 1


def test_fixed_expense_flow(logged_in, db):
    c = logged_in
    due = today() + timedelta(days=2)
    r = c.post("/sabit", data={"ad": "Kamyon leasing", "tutar": "18.500", "ilk_vade": due.isoformat(),
                               "taksit": "24", "hatirlat": "3", "kategori_id": str(_cat(db, "Leasing"))},
               follow_redirects=False)
    assert r.status_code == 303
    fx = db.query(FixedExpense).one()
    assert fx.amount == 18_500_00 and fx.total_count == 24
    inst = db.query(Installment).order_by(Installment.due_date).first()
    assert inst.due_date == due

    dash = c.get("/").text
    assert "Kamyon leasing" in dash and "2 gün kaldı" in dash          # panoda uyarı

    fatma = db.query(User).filter_by(name="Fatma").one()
    c.post(f"/taksit/{inst.id}/ode", data={"gun": today().isoformat(), "tutar": "18.500,00",
                                           "kim": str(fatma.id), "donus": "/"})
    db.refresh(inst)
    assert inst.is_paid
    tx = db.query(Transaction).one()
    assert tx.amount == 18_500_00 and tx.partner_id == fatma.id
    assert "2 gün kaldı" not in c.get("/").text
    # aynı taksit ikinci kez ödenemez
    c.post(f"/taksit/{inst.id}/ode", data={"gun": today().isoformat(), "tutar": "18.500", "kim": "ortak"})
    assert db.query(Transaction).count() == 1

    # ödeme kaydı silinince taksit yeniden açılır
    c.post(f"/kayitlar/{tx.id}/sil")
    db.refresh(inst)
    assert not inst.is_paid and "2 gün kaldı" in c.get("/").text

    # tutar değişince ödenmemiş taksitler güncellenir
    c.post(f"/sabit/{fx.id}/guncelle", data={"ad": "Kamyon leasing", "tutar": "19.000", "hatirlat": "5"})
    db.refresh(inst)
    assert inst.amount == 19_000_00

    c.post(f"/sabit/{fx.id}/sonlandir")
    assert db.query(FixedExpense).count() == 0 and db.query(Installment).count() == 0


def test_partner_balance_and_settlement_pages(logged_in, db):
    c = logged_in
    fatma = db.query(User).filter_by(name="Fatma").one()
    kerem = db.query(User).filter_by(name="Kerem").one()
    c.post("/kayitlar/yeni", data={"tur": "gider", "gun": today().isoformat(), "tutar": "1000",
                                   "kategori_id": str(_cat(db, "Yakıt")), "kim": str(fatma.id)})
    page = c.get("/ortaklar").text
    assert "<b>Kerem</b>, <b>Fatma</b> ortağına" in page and "500,00 ₺" in page
    c.post("/ortaklar/hesaplasma", data={"gun": today().isoformat(), "kimden": str(kerem.id),
                                         "kime": str(fatma.id), "tutar": "500"})
    assert db.query(Settlement).count() == 1
    assert "Hesap denk" in c.get("/ortaklar").text
    # kendine ödeme kaydedilemez
    c.post("/ortaklar/hesaplasma", data={"gun": today().isoformat(), "kimden": str(kerem.id),
                                         "kime": str(kerem.id), "tutar": "5"})
    assert db.query(Settlement).count() == 1


def test_shares_must_total_100(logged_in, db):
    c = logged_in
    a, b = db.query(User).order_by(User.id).all()
    c.post("/ayarlar/pay", data={f"pay_{a.id}": "60", f"pay_{b.id}": "30"})
    db.refresh(a)
    assert a.share_bp == 5000
    c.post("/ayarlar/pay", data={f"pay_{a.id}": "62,5", f"pay_{b.id}": "37,5"})
    db.refresh(a); db.refresh(b)
    assert (a.share_bp, b.share_bp) == (6250, 3750)


def test_report_and_exports(logged_in, db):
    c = logged_in
    t = today()
    fatma = db.query(User).filter_by(name="Fatma").one()
    rows = [("gider", "1.200,50", "Yakıt", str(fatma.id), "=HYPERLINK(\"http://x\")"),
            ("gider", "800", "Yemek", "ortak", "öğle"),
            ("gelir", "10.000", "İş geliri", "ortak", "hakediş")]
    for tur, tutar, kat, kim, acik in rows:
        c.post("/kayitlar/yeni", data={"tur": tur, "gun": t.isoformat(), "tutar": tutar,
                                       "kategori_id": str(_cat(db, kat)), "kim": kim, "aciklama": acik})
    page = c.get("/rapor").text
    assert "10.000,00 ₺" in page and "2.000,50 ₺" in page and "7.999,50 ₺" in page

    start, end = t.replace(day=1).isoformat(), t.isoformat()
    r = c.get(f"/rapor/excel?bas={start}&bit={end}")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Özet", "Giderler", "Gelirler", "Sabit ödemeler", "Hesaplaşmalar"]
    ozet = {row[0]: row[1] for row in wb["Özet"].iter_rows(values_only=True) if row[0]}
    assert ozet["Toplam gelir"] == 10000 and ozet["Toplam gider"] == 2000.5
    assert ozet["Net (gelir − gider)"] == 7999.5
    giderler = list(wb["Giderler"].iter_rows(values_only=True))
    assert giderler[1][4] == 1200.5 and giderler[3][4] == 2000.5       # satır + toplam
    # kullanıcı metni formül olarak çalışmamalı
    assert wb["Giderler"]["C2"].data_type == "s"

    r = c.get(f"/rapor/csv?bas={start}&bit={end}")
    text = r.content.decode("utf-8-sig")
    lines = text.strip().split("\r\n")
    assert lines[0].startswith("Tarih;Tür;Kategori") and len(lines) == 4
    assert "1200,50" in text and "'=HYPERLINK" in text


def test_category_management(logged_in, db):
    c = logged_in
    c.post("/ayarlar/kategori", data={"ad": "Otoyol", "tur": "gider"})
    cat = db.query(Category).filter_by(name="Otoyol").one()
    c.post("/ayarlar/kategori", data={"ad": "otoyol", "tur": "gider"})           # aynısı ikinci kez eklenmez
    assert db.query(Category).filter(Category.name.ilike("otoyol")).count() == 1
    c.post("/kayitlar/yeni", data={"tur": "gider", "gun": today().isoformat(), "tutar": "90",
                                   "kategori_id": str(cat.id), "kim": "ortak"})
    c.post(f"/ayarlar/kategori/{cat.id}/durum")                                  # kullanımda: silinmez, kapanır
    db.refresh(cat)
    assert cat.is_active is False
    assert "Otoyol" not in c.get("/kayitlar/yeni?tur=gider").text
    assert "Otoyol" in c.get("/kayitlar").text                                   # eski kayıt duruyor


def test_mistyped_year_is_rejected_everywhere(logged_in, db):
    c = logged_in
    fatma = db.query(User).filter_by(name="Fatma").one()
    kerem = db.query(User).filter_by(name="Kerem").one()
    r = c.post("/kayitlar/yeni", data={"tur": "gider", "gun": "0026-10-08", "tutar": "5.000",
                                       "kategori_id": str(_cat(db, "Yakıt")), "kim": str(fatma.id)})
    assert r.status_code == 422 and db.query(Transaction).count() == 0
    c.post("/ortaklar/hesaplasma", data={"gun": "0026-10-08", "kimden": str(kerem.id),
                                         "kime": str(fatma.id), "tutar": "10"})
    assert db.query(Settlement).count() == 0


def test_odd_input_never_causes_server_error(logged_in, db):
    c = logged_in
    base = {"tur": "gider", "gun": today().isoformat(), "kategori_id": str(_cat(db, "Yemek")), "kim": "ortak"}
    for tutar in ["1,5.", "12,.", "9" * 5000, "²"]:
        assert c.post("/kayitlar/yeni", data={**base, "tutar": tutar}).status_code == 422
    big = "9" * 25
    assert c.post("/kayitlar/yeni", data={**base, "tutar": "5", "kim": big}).status_code == 422
    assert c.post("/kayitlar/yeni", data={**base, "tutar": "5", "kategori_id": "²"}).status_code == 422
    for url in [f"/kayitlar/{big}", f"/kayitlar?kategori={big}&kim={big}", f"/fisler/{big}/dosya",
                "/kayitlar?ay=0026-10", "/rapor?ay=abc", "/rapor/excel?bas=x&bit=y"]:
        assert c.get(url).status_code in (200, 404), url
    a, b = db.query(User).order_by(User.id).all()
    for val in ["inf", "1e400", "nan", "-5"]:
        assert c.post("/ayarlar/pay", data={f"pay_{a.id}": val, f"pay_{b.id}": "50"}).status_code == 200
    db.refresh(a)
    assert a.share_bp == 5000


def test_share_field_shows_what_was_saved(logged_in, db):
    c = logged_in
    a, b = db.query(User).order_by(User.id).all()
    c.post("/ayarlar/pay", data={f"pay_{a.id}": "33,05", f"pay_{b.id}": "66,95"})
    page = c.get("/ayarlar").text
    assert 'value="33,05"' in page and 'value="66,95"' in page


def test_turkish_category_names_do_not_duplicate(logged_in, db):
    c = logged_in
    assert c.post("/ayarlar/kategori", data={"ad": "Şoför", "tur": "gider"}).status_code == 200
    assert c.post("/ayarlar/kategori", data={"ad": "ŞOFÖR", "tur": "gider"}).status_code == 200
    assert c.post("/ayarlar/kategori", data={"ad": "iş geliri", "tur": "gelir"}).status_code == 200
    assert db.query(Category).filter(Category.name.in_(["Şoför", "ŞOFÖR"])).count() == 1
    assert db.query(Category).filter_by(kind="gelir").count() == 2
    other = db.query(Category).filter_by(name="Yemek").one()
    assert c.post(f"/ayarlar/kategori/{other.id}/ad", data={"ad": "şoför"}).status_code == 200
    db.refresh(other)
    assert other.name == "Yemek"


def test_password_change_invalidates_old_sessions(logged_in, db):
    c = logged_in
    old_cookie = c.cookies.get("defter_oturum")
    c.post("/ayarlar/sifre", data={"eski": "gizli-sifre-1", "yeni": "yepyeni-sifre", "tekrar": "yepyeni-sifre"})
    assert c.get("/ayarlar").status_code == 200                      # bu cihaz açık kalır
    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app, base_url="http://testserver", cookies={"defter_oturum": old_cookie}) as stale:
        assert stale.get("/ayarlar", follow_redirects=False).headers["location"] == "/giris"


def test_anonymous_post_is_turned_away_before_upload(client, db):
    client.post("/kurulum", data={"ad1": "F", "eposta1": "f@x.co", "sifre1": "gizli-sifre-1",
                                  "ad2": "K", "eposta2": "k@x.co", "sifre2": "gizli-sifre-2"})
    client.post("/cikis")
    r = client.post("/kayitlar/yeni", data={"tutar": "5"},
                    files=[("fisler", ("f.png", _png(size=(20, 20)), "image/png"))], follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/giris"
    assert db.query(Transaction).count() == 0


def test_control_characters_do_not_break_export_or_mail(logged_in, db):
    c = logged_in
    c.post("/kayitlar/yeni", data={"tur": "gider", "gun": today().isoformat(), "tutar": "10",
                                   "kategori_id": str(_cat(db, "Yemek")), "kim": "ortak",
                                   "aciklama": "abc\x0bdef\r\nikinci satır"})
    tx = db.query(Transaction).one()
    assert tx.description == "abc def ikinci satır"
    # eski/bozuk veri olsa bile döküm çalışmalı
    tx.description = "x\x0by"
    db.commit()
    assert c.get(f"/rapor/excel?bas={today().isoformat()}&bit={today().isoformat()}").status_code == 200
    c.post("/sabit", data={"ad": "Kira\r\nBcc: x@y.z", "tutar": "100", "ilk_vade": today().isoformat(),
                           "kategori_id": str(_cat(db, "Ek gider"))})
    assert db.query(FixedExpense).one().name == "Kira Bcc: x@y.z"


def test_deleting_payment_of_ended_fixed_expense_leaves_no_hidden_debt(logged_in, db):
    c = logged_in
    c.post("/sabit", data={"ad": "Kira", "tutar": "100", "ilk_vade": today().isoformat(),
                           "kategori_id": str(_cat(db, "Ek gider"))})
    fx = db.query(FixedExpense).one()
    inst = db.query(Installment).order_by(Installment.due_date).first()
    c.post(f"/taksit/{inst.id}/ode", data={"gun": today().isoformat(), "tutar": "100", "kim": "ortak"})
    c.post(f"/sabit/{fx.id}/sonlandir")
    db.expire_all()
    assert db.query(Installment).count() == 1                      # yalnızca ödenmiş olan kaldı
    tx = db.query(Transaction).one()
    c.post(f"/kayitlar/{tx.id}/sil")
    assert db.query(Installment).count() == 0 and db.query(Transaction).count() == 0


# ---------------------------------------------------------------- yeni tasarım, simgeler, telefona ekleme

def test_install_files_are_public_and_valid(client, db):
    m = client.get("/manifest.webmanifest")
    assert m.status_code == 200 and m.json()["display"] == "standalone"
    for icon in m.json()["icons"]:
        assert client.get(icon["src"]).status_code == 200, icon
    sw = client.get("/sw.js")
    assert sw.status_code == 200 and "javascript" in sw.headers["content-type"] and "/cevrimdisi" in sw.text
    off = client.get("/cevrimdisi")
    assert off.status_code == 200 and "Bağlantı yok" in off.text
    for path in ["/static/brand/apple-touch-icon.png", "/static/brand/favicon.png", "/static/js/theme.js",
                 "/static/fonts/archivo-latin-wdth-normal.woff2", "/static/fonts/archivo-latin-ext-wdth-normal.woff2"]:
        assert client.get(path).status_code == 200, path


def test_setup_key_is_required_when_configured(client, db, monkeypatch):
    from app.config import settings as cfg
    monkeypatch.setattr(cfg, "setup_key", "cok-gizli-anahtar")
    data = {"ad1": "F", "eposta1": "f@x.co", "sifre1": "gizli-sifre-1",
            "ad2": "K", "eposta2": "k@x.co", "sifre2": "gizli-sifre-2"}
    assert "Kurulum anahtarı" in client.get("/kurulum").text
    assert client.post("/kurulum", data=data).status_code == 422
    assert client.post("/kurulum", data={**data, "anahtar": "yanlis"}).status_code == 422
    assert db.query(User).count() == 0
    r = client.post("/kurulum", data={**data, "anahtar": "cok-gizli-anahtar"}, follow_redirects=False)
    assert r.status_code == 303 and db.query(User).count() == 2


def test_categories_get_icons_and_icon_can_change(logged_in, db):
    c = logged_in
    icons = {cat.name: cat.icon for cat in db.query(Category)}
    assert icons["Yakıt"] == "yakit" and icons["Yevmiye"] == "isci" and icons["Yemek"] == "yemek"
    assert icons["Bakım / Onarım"] == "anahtar" and icons["Leasing"] == "belge" and icons["İş geliri"] == "canta"
    c.post("/ayarlar/kategori", data={"ad": "Otoyol / köprü", "tur": "gider"})
    c.post("/ayarlar/kategori", data={"ad": "Bilinmeyen şey", "tur": "gider"})
    assert db.query(Category).filter_by(name="Otoyol / köprü").one().icon == "yol"
    other = db.query(Category).filter_by(name="Bilinmeyen şey").one()
    assert other.icon == "etiket"
    c.post(f"/ayarlar/kategori/{other.id}/ad", data={"ad": "Bilinmeyen şey", "simge": "kalkan"})
    c.post(f"/ayarlar/kategori/{other.id}/ad", data={"ad": "Bilinmeyen şey", "simge": "<script>"})   # yok sayılır
    db.refresh(other)
    assert other.icon == "kalkan"


def test_old_database_gets_icon_column(db):
    from sqlalchemy import inspect, text
    from app.db import _upgrade_schema, engine
    with engine.begin() as conn:
        conn.execute(text("DROP TABLE categories"))
        conn.execute(text("CREATE TABLE categories (id INTEGER PRIMARY KEY, name VARCHAR(60), kind VARCHAR(10), "
                          "sort INTEGER, is_active BOOLEAN)"))
        conn.execute(text("INSERT INTO categories (name, kind, sort, is_active) VALUES "
                          "('Yakıt','gider',0,1), ('Kasko','gider',1,1), ('İş geliri','gelir',0,1)"))
    _upgrade_schema()
    assert "icon" in {col["name"] for col in inspect(engine).get_columns("categories")}
    with engine.connect() as conn:
        got = dict(conn.execute(text("SELECT name, icon FROM categories")).all())
    assert got == {"Yakıt": "yakit", "Kasko": "kalkan", "İş geliri": "canta"}
    _upgrade_schema()   # ikinci kez çalışınca sorun çıkmaz


def test_dashboard_summary_and_charts(logged_in, db):
    c = logged_in
    empty = c.get("/").text
    assert "Bu ay için kayıt yok" in empty and "data-chart" not in empty
    t = today()
    for tutar, kat, tur in [("1.000", "Yakıt", "gider"), ("250,50", "Yemek", "gider"), ("4.000", "İş geliri", "gelir")]:
        c.post("/kayitlar/yeni", data={"tur": tur, "gun": t.isoformat(), "tutar": tutar,
                                       "kategori_id": str(_cat(db, kat)), "kim": "ortak"})
    page = c.get("/").text
    assert "2.749,50 ₺" in page                      # net
    assert "4.000,00 ₺" in page and "1.250,50 ₺" in page
    assert page.count('class="col"') == 14           # 14 günlük grafik
    assert "Bu ay gider nereye gitti?" in page and "%80" in page and "%20" in page


def test_daily_series_and_top_categories(db):
    from datetime import date
    from app.services import reports
    from app.services.reports import CategoryRow
    cat = db.query(Category).filter_by(name="Yakıt").one()
    db.add_all([
        Transaction(kind="gider", day=date(2026, 10, 1), amount=500, category_id=cat.id),
        Transaction(kind="gider", day=date(2026, 9, 30), amount=300, category_id=cat.id),
        Transaction(kind="gelir", day=date(2026, 10, 1), amount=900, category_id=cat.id),
        Transaction(kind="gider", day=date(2026, 9, 10), amount=999, category_id=cat.id),   # aralık dışı
    ])
    db.commit()
    series = reports.daily_series(db, date(2026, 10, 2), 14)
    assert len(series) == 14 and series[0].day == date(2026, 9, 19) and series[-1].day == date(2026, 10, 2)
    by_day = {p.day: (p.income, p.expense) for p in series}
    assert by_day[date(2026, 10, 1)] == (900, 500) and by_day[date(2026, 9, 30)] == (0, 300)
    assert sum(p.expense for p in series) == 800

    rows = [CategoryRow(f"K{i}", 100 - i, 1, 10.0) for i in range(9)]
    top = reports.top_categories(rows, 5)
    assert len(top) == 6 and top[-1].name == "Diğer 4 kategori"
    assert sum(r.amount for r in top) == sum(r.amount for r in rows)
    assert reports.top_categories(rows[:6], 5) == rows[:6]   # tek satır için "Diğer" açılmaz


def test_large_jpeg_is_downscaled_with_correct_shape_and_orientation(logged_in, db):
    import numpy as np
    c = logged_in
    # 4032x3024 (telefon fotoğrafı ölçüsü), EXIF yönü 6: gösterimde dik durmalı
    arr = np.zeros((3024, 4032, 3), dtype=np.uint8)
    arr[:, :2016] = (220, 30, 30)            # sol yarı kırmızı
    arr[:, 2016:] = (30, 30, 220)            # sağ yarı mavi
    img = Image.fromarray(arr)
    exif = img.getexif()
    exif[0x0112] = 6
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=90, exif=exif)
    r = c.post("/kayitlar/yeni", data={"tur": "gider", "gun": today().isoformat(), "tutar": "10",
                                       "kategori_id": str(_cat(db, "Yemek")), "kim": "ortak"},
               files=[("fisler", ("foto.jpg", buf.getvalue(), "image/jpeg"))])
    assert r.status_code == 200
    rec = db.query(Receipt).one()
    out = Image.open(settings.upload_dir / rec.path)
    assert out.size == (1500, 2000)                      # döndürülmüş, uzun kenar 2000
    top, bottom = out.getpixel((750, 200)), out.getpixel((750, 1800))
    assert top[0] > 150 and top[2] < 100                 # üstte kırmızı
    assert bottom[2] > 150 and bottom[0] < 100           # altta mavi
    assert not out.getexif().get(0x0112)                 # yön bilgisi artık gerekmiyor, EXIF atıldı
