"""Muhasebeciye verilecek döküm: Excel (çok sayfalı) ve CSV."""
from __future__ import annotations

import csv
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import settings
from ..models import GELIR, GIDER, PAYMENT_LABELS, FixedExpense, Installment, PartnerPayment, Transaction
from ..timeutil import fmt_date
from ..money import TRY, format_money
from . import balance, rates, reports

TL = '#,##0.00 "₺"'
DATE = "DD.MM.YYYY"
HEAD_FILL = PatternFill("solid", start_color="16233B")
HEAD_FONT = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="D9DEE3")


def _safe(text: str | None) -> str:
    """CSV için: hücrenin formül olarak çalışmasını engeller (=, +, -, @ ile başlayanlar)."""
    text = text or ""
    return "'" + text if text[:1] in ("=", "+", "-", "@", "\t", "\r") else text


def _defuse_formulas(wb: Workbook) -> None:
    """Excel için: kullanıcı metni '=' ile başlıyorsa formül değil düz metin olarak kalsın.

    Dosyada bilerek yazılmış formül yok; formül görünen her hücre kullanıcı metnidir.
    """
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"


def _transactions(db: Session, start: date, end: date, kind: str | None = None):
    q = select(Transaction).where(Transaction.day.between(start, end))
    if kind:
        q = q.where(Transaction.kind == kind)
    return db.scalars(q.order_by(Transaction.day, Transaction.id)).unique().all()


def _t(text: str | None) -> str:
    """Excel'in kabul etmediği denetim karakterlerini atar."""
    return ILLEGAL_CHARACTERS_RE.sub("", text or "")


def _header(ws, row: int, titles: list[str]) -> None:
    for col, title in enumerate(titles, 1):
        c = ws.cell(row=row, column=col, value=title)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 20


def _widths(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _tx_sheet(ws, rows: list[Transaction], who_title: str) -> None:
    _header(ws, 1, ["Tarih", "Kategori", "Açıklama", who_title, "Tutar", "Fiş sayısı", "Giren", "Kayıt no"])
    for r, tx in enumerate(rows, 2):
        ws.cell(r, 1, tx.day).number_format = DATE
        ws.cell(r, 2, _t(tx.category.name if tx.category else "Kategorisiz"))
        ws.cell(r, 3, _t(tx.description))
        ws.cell(r, 4, _t(tx.partner.name if tx.partner else "Ortak hesap"))
        ws.cell(r, 5, tx.amount / 100).number_format = TL
        ws.cell(r, 6, len(tx.receipts))
        ws.cell(r, 7, _t(tx.created_by.name if tx.created_by else ""))
        ws.cell(r, 8, tx.id)
    last = len(rows) + 1
    total_row = last + 1
    ws.cell(total_row, 4, "Toplam").font = BOLD
    c = ws.cell(total_row, 5, sum(tx.amount for tx in rows) / 100)
    c.number_format, c.font = TL, BOLD
    c.border = Border(top=THIN)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:H{last}"
    _widths(ws, [13, 20, 44, 18, 16, 11, 16, 10])


def build_xlsx(db: Session, start: date, end: date) -> bytes:
    rep = reports.period_report(db, start, end)
    expenses = _transactions(db, start, end, GIDER)
    incomes = _transactions(db, start, end, GELIR)

    wb = Workbook()

    # --- Özet ---
    ws = wb.active
    ws.title = "Özet"
    ws["A1"] = _t(settings.app_name)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Dönem: {fmt_date(start)} – {fmt_date(end)}"
    row = 4
    for label, value, bold in [
        ("Toplam gelir", rep.income, False),
        ("Toplam gider", rep.expense, False),
        ("Net (gelir − gider)", rep.net, True),
        ("Dönemde vadesi gelen, ödenmemiş sabit giderler", rep.unpaid_fixed, False),
        ("Onlar da ödenince kalan", rep.net_after_fixed, True),
    ]:
        ws.cell(row, 1, label).font = BOLD if bold else Font()
        c = ws.cell(row, 2, value / 100)
        c.number_format = TL
        c.font = BOLD if bold else Font()
        row += 1

    row += 1
    _header(ws, row, ["Gider kategorisi", "Tutar", "Pay", "Kayıt"])
    for cat in rep.expense_categories:
        row += 1
        ws.cell(row, 1, _t(cat.name))
        ws.cell(row, 2, cat.amount / 100).number_format = TL
        ws.cell(row, 3, cat.pct / 100).number_format = "0.0%"
        ws.cell(row, 4, cat.count)

    row += 2
    _header(ws, row, ["Gelir kategorisi", "Tutar", "Pay", "Kayıt"])
    for cat in rep.income_categories:
        row += 1
        ws.cell(row, 1, _t(cat.name))
        ws.cell(row, 2, cat.amount / 100).number_format = TL
        ws.cell(row, 3, cat.pct / 100).number_format = "0.0%"
        ws.cell(row, 4, cat.count)

    row += 2
    _header(ws, row, ["Dönemde kim ödedi / kim aldı", "Ödediği gider", "Aldığı gelir"])
    for p in rep.partners:
        row += 1
        ws.cell(row, 1, _t(p.name))
        ws.cell(row, 2, p.paid / 100).number_format = TL
        ws.cell(row, 3, p.collected / 100).number_format = TL

    sheet = balance.compute(db)
    row += 2
    _header(ws, row, ["Ortakların şirketle hesabı (bugüne kadar, tüm kayıtlar)", "Bakiye", "Durum"])
    for line in sheet.lines:
        row += 1
        ws.cell(row, 1, _t(line.user.name))
        ws.cell(row, 2, line.balance / 100).number_format = TL
        ws.cell(row, 3, "Şirketten alacaklı" if line.balance > 0 else "Şirkete borçlu" if line.balance < 0 else "Denk")
    _widths(ws, [48, 18, 16, 10])

    # --- Kayıtlar ---
    _tx_sheet(wb.create_sheet("Giderler"), expenses, "Ödeyen")
    _tx_sheet(wb.create_sheet("Gelirler"), incomes, "Teslim alan")

    # --- Sabit ödemeler ---
    ws = wb.create_sheet("Sabit ödemeler")
    _header(ws, 1, ["Sabit gider", "Taksit", "Vade", "Tutar (₺)", "Döviz tutarı", "Durum", "Ödeme tarihi", "Ödeyen"])
    rate = rates.latest(db)
    insts = db.scalars(
        select(Installment)
        .join(FixedExpense)
        .where(Installment.due_date.between(start, end))
        .order_by(Installment.due_date, Installment.id)
    ).all()
    for r, inst in enumerate(insts, 2):
        tx = inst.transaction
        ws.cell(r, 1, _t(inst.fixed.name))
        ws.cell(r, 2, f"{inst.seq}/{inst.fixed.total_count}" if inst.fixed.total_count else str(inst.seq))
        ws.cell(r, 3, inst.due_date).number_format = DATE
        cur = inst.fixed.currency
        if tx:
            try_amount = tx.amount
        elif cur == TRY:
            try_amount = inst.amount
        else:
            try_amount = rates.convert(inst.amount, rate.value) if rate else None
        if try_amount is not None:
            ws.cell(r, 4, try_amount / 100).number_format = TL
        if cur != TRY:
            ws.cell(r, 5, format_money(inst.amount, cur))
        ws.cell(r, 6, "Ödendi" if tx else "Ödenmedi")
        if tx:
            ws.cell(r, 7, tx.day).number_format = DATE
            ws.cell(r, 8, _t(tx.partner.name if tx.partner else "Ortak hesap"))
    ws.freeze_panes = "A2"
    _widths(ws, [30, 10, 13, 16, 16, 12, 14, 18])

    ws = wb.create_sheet("Ortak ödemeleri")
    _header(ws, 1, ["Tarih", "Ortak", "İşlem", "Tutar", "Not"])
    pays = db.scalars(
        select(PartnerPayment).where(PartnerPayment.day.between(start, end)).order_by(PartnerPayment.day)
    ).all()
    for r, p in enumerate(pays, 2):
        ws.cell(r, 1, p.day).number_format = DATE
        ws.cell(r, 2, _t(p.user.name))
        ws.cell(r, 3, PAYMENT_LABELS.get(p.direction, p.direction))
        ws.cell(r, 4, p.amount / 100).number_format = TL
        ws.cell(r, 5, _t(p.note))
    ws.freeze_panes = "A2"
    _widths(ws, [13, 20, 24, 16, 40])

    _defuse_formulas(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_csv(db: Session, start: date, end: date) -> bytes:
    """Türkçe Excel'in doğrudan açabildiği biçim: noktalı virgül ayraç, UTF-8 BOM."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(["Tarih", "Tür", "Kategori", "Açıklama", "Ödeyen / teslim alan", "Tutar", "Fiş sayısı", "Kayıt no"])
    for tx in _transactions(db, start, end):
        w.writerow([
            tx.day.strftime("%d.%m.%Y"),
            "Gelir" if tx.kind == GELIR else "Gider",
            _safe(tx.category.name if tx.category else "Kategorisiz"),
            _safe(tx.description),
            _safe(tx.partner.name if tx.partner else "Ortak hesap"),
            f"{tx.amount / 100:.2f}".replace(".", ","),
            len(tx.receipts),
            tx.id,
        ])
    return ("﻿" + buf.getvalue()).encode("utf-8")
